from __future__ import annotations

import copy
import os
import re
import shutil
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Inches


DOCX_PATH = Path(r"D:\AAA\备份论文.docx")
BACKUP_PATH = Path(r"D:\AAA\备份论文_插图前备份.docx")
PICTURES_DIR = Path(r"C:\Users\王杰\Desktop\论文\pictures")
REVISED_DIR = PICTURES_DIR / "修订图"


@dataclass
class FigureSpec:
    anchor_prefix: str
    figure_label: str
    title: str
    image_path: Path
    width_inches: float
    intro_text: str


def normalize(text: str) -> str:
    return re.sub(r"\s+", "", text or "")


def make_intro(label: str, title: str, purpose: str) -> str:
    return f"{purpose}，如{label} {title}所示。"


FIGURES: list[FigureSpec] = [
    FigureSpec(
        anchor_prefix="1.1",
        figure_label="图1-1",
        title="系统应用场景示意图",
        image_path=PICTURES_DIR / "图1-1-系统应用场景示意图.png",
        width_inches=6.0,
        intro_text=make_intro("图1-1", "系统应用场景示意图", "为直观展示本课题面向的灾后临时安置点应用环境以及系统服务对象之间的关系，本节给出系统应用场景示意"),
    ),
    FigureSpec(
        anchor_prefix="2.1",
        figure_label="图2-1",
        title="系统业务流程图",
        image_path=PICTURES_DIR / "图2-1-系统业务流程图.png",
        width_inches=6.0,
        intro_text=make_intro("图2-1", "系统业务流程图", "为说明系统从现场采集、通信传输到终端展示与响应处理的完整业务过程，本节给出系统业务流程"),
    ),
    FigureSpec(
        anchor_prefix="3.3.1",
        figure_label="图3-1",
        title="系统总体架构图",
        image_path=REVISED_DIR / "图3-1-系统总体架构图.png",
        width_inches=6.2,
        intro_text=make_intro("图3-1", "系统总体架构图", "为了说明系统各层次模块之间的组织关系以及多端协同方式，本节给出系统总体架构"),
    ),
    FigureSpec(
        anchor_prefix="3.2",
        figure_label="图3-2",
        title="系统硬件连接图",
        image_path=PICTURES_DIR / "图3-2-硬件连接图.png",
        width_inches=6.2,
        intro_text=make_intro("图3-2", "系统硬件连接图", "为了展示 STM32 主控与各类传感器、执行器及 ESP32 网关之间的实际引脚连接关系，本节给出系统硬件连接图"),
    ),
    FigureSpec(
        anchor_prefix="3.2",
        figure_label="图3-3",
        title="系统供电方案图",
        image_path=PICTURES_DIR / "图3-3-供电方案图.png",
        width_inches=3.1,
        intro_text=make_intro("图3-3", "系统供电方案图", "为了说明系统各模块的供电电压来源及共地关系，本节给出系统供电方案"),
    ),
    FigureSpec(
        anchor_prefix="3.2.3",
        figure_label="图3-4",
        title="系统硬件实物图",
        image_path=PICTURES_DIR / "硬件实物图.jpg",
        width_inches=5.2,
        intro_text=make_intro("图3-4", "系统硬件实物图", "为了反映本课题硬件平台的实际搭建情况与模块集成效果，本节给出系统硬件实物图"),
    ),
    FigureSpec(
        anchor_prefix="3.3.3",
        figure_label="图3-5",
        title="系统数据流图",
        image_path=PICTURES_DIR / "图3-3-数据流图.png",
        width_inches=6.2,
        intro_text=make_intro("图3-5", "系统数据流图", "为了说明系统上行监测数据与下行控制信息在各节点之间的传递路径，本节给出系统数据流图"),
    ),
    FigureSpec(
        anchor_prefix="3.7.1",
        figure_label="图3-6",
        title="数据库ER图",
        image_path=PICTURES_DIR / "图3-4-数据库ER图.png",
        width_inches=6.0,
        intro_text=make_intro("图3-6", "数据库ER图", "为了说明系统本地数据库中各核心数据表的组成以及业务关联关系，本节给出数据库 ER 图"),
    ),
    FigureSpec(
        anchor_prefix="4.1.2",
        figure_label="图4-1",
        title="GP2Y10脉冲驱动时序图",
        image_path=PICTURES_DIR / "图4-2-GP2Y10脉冲驱动时序图.png",
        width_inches=6.0,
        intro_text=make_intro("图4-1", "GP2Y10脉冲驱动时序图", "为了说明 GP2Y10 粉尘传感器在 LED 驱动与 ADC 采样过程中的关键时序关系，本节给出 GP2Y10 脉冲驱动时序图"),
    ),
    FigureSpec(
        anchor_prefix="4.1.3",
        figure_label="图4-2",
        title="水电剩余量模拟原理图",
        image_path=PICTURES_DIR / "图4-6-水电剩余量模拟原理图.png",
        width_inches=5.0,
        intro_text=make_intro("图4-2", "水电剩余量模拟原理图", "为了说明系统对电量消耗与用水累计进行估算和剩余量换算的实现思路，本节给出水电剩余量模拟原理图"),
    ),
    FigureSpec(
        anchor_prefix="4.2.1",
        figure_label="图4-3",
        title="STM32主程序流程图",
        image_path=PICTURES_DIR / "图4-1-STM32主程序流程图.png",
        width_inches=6.0,
        intro_text=make_intro("图4-3", "STM32主程序流程图", "为了说明 STM32 端主循环中各项任务的初始化顺序和运行流程，本节给出 STM32 主程序流程图"),
    ),
    FigureSpec(
        anchor_prefix="4.2.3",
        figure_label="图4-4",
        title="三级报警机制示意图",
        image_path=PICTURES_DIR / "图4-3-三级报警机制示意图.png",
        width_inches=5.4,
        intro_text=make_intro("图4-4", "三级报警机制示意图", "为了展示系统对各监测指标进行正常、预警和告警三级判定的逻辑过程，本节给出三级报警机制示意图"),
    ),
    FigureSpec(
        anchor_prefix="4.2.3",
        figure_label="图4-5",
        title="滞回防抖示意图",
        image_path=PICTURES_DIR / "图4-5-滞回防抖示意图.png",
        width_inches=5.4,
        intro_text=make_intro("图4-5", "滞回防抖示意图", "为了说明系统在阈值临界附近通过滞回与防抖机制抑制频繁误触发的实现方式，本节给出滞回防抖示意图"),
    ),
    FigureSpec(
        anchor_prefix="4.3.3",
        figure_label="图4-6",
        title="MQTT通信时序图",
        image_path=PICTURES_DIR / "图4-2-MQTT通信时序图.png",
        width_inches=6.0,
        intro_text=make_intro("图4-6", "MQTT通信时序图", "为了说明 STM32、ESP32、MQTT 平台以及 Qt 与 Web 端之间的消息交互顺序，本节给出 MQTT 通信时序图"),
    ),
    FigureSpec(
        anchor_prefix="4.4.2",
        figure_label="图4-7",
        title="Qt实时监控页面",
        image_path=PICTURES_DIR / "qt实时监控页面.png",
        width_inches=5.8,
        intro_text=make_intro("图4-7", "Qt实时监控页面", "为了展示 Qt 客户端对环境与资源监测数据的实时展示效果，本节给出 Qt 实时监控页面"),
    ),
    FigureSpec(
        anchor_prefix="4.4.3",
        figure_label="图4-8",
        title="Qt水电管理页面",
        image_path=PICTURES_DIR / "qt水电管理页面.png",
        width_inches=5.8,
        intro_text=make_intro("图4-8", "Qt水电管理页面", "为了展示客户端对水电资源状态、阈值设置和相关控制功能的界面实现效果，本节给出 Qt 水电管理页面"),
    ),
    FigureSpec(
        anchor_prefix="4.4.4",
        figure_label="图4-9",
        title="Qt设备管理页面",
        image_path=PICTURES_DIR / "qt设备管理页面.png",
        width_inches=5.8,
        intro_text=make_intro("图4-9", "Qt设备管理页面", "为了展示客户端对设备台账、运行状态与管理信息的统一展示效果，本节给出 Qt 设备管理页面"),
    ),
    FigureSpec(
        anchor_prefix="4.5.2",
        figure_label="图4-10",
        title="Web看板页面",
        image_path=PICTURES_DIR / "web页面.png",
        width_inches=6.0,
        intro_text=make_intro("图4-10", "Web看板页面", "为了展示浏览器端对环境状态、资源信息和告警内容的可视化呈现效果，本节给出 Web 看板页面"),
    ),
    FigureSpec(
        anchor_prefix="4.5.4",
        figure_label="图4-11",
        title="Web一键求助提示界面",
        image_path=PICTURES_DIR / "webtoast弹窗.png",
        width_inches=3.3,
        intro_text=make_intro("图4-11", "Web一键求助提示界面", "为了展示 Web 端一键求助功能触发后的界面反馈效果，本节给出 Web 一键求助提示界面"),
    ),
    FigureSpec(
        anchor_prefix="5.6.1",
        figure_label="图5-1",
        title="72小时连续运行数据曲线",
        image_path=PICTURES_DIR / "图5-2-72小时连续运行数据曲线.png",
        width_inches=5.8,
        intro_text=make_intro("图5-1", "72小时连续运行数据曲线", "为了验证系统在连续运行条件下的数据稳定性与整体运行可靠性，本节给出 72 小时连续运行数据曲线"),
    ),
]


def clone_paragraph_format(src, dst):
    dst.style = src.style
    dst.paragraph_format.left_indent = src.paragraph_format.left_indent
    dst.paragraph_format.right_indent = src.paragraph_format.right_indent
    dst.paragraph_format.first_line_indent = src.paragraph_format.first_line_indent
    dst.paragraph_format.space_before = src.paragraph_format.space_before
    dst.paragraph_format.space_after = src.paragraph_format.space_after
    dst.paragraph_format.line_spacing = src.paragraph_format.line_spacing
    dst.paragraph_format.line_spacing_rule = src.paragraph_format.line_spacing_rule
    dst.paragraph_format.keep_together = src.paragraph_format.keep_together
    dst.paragraph_format.keep_with_next = src.paragraph_format.keep_with_next
    dst.paragraph_format.page_break_before = src.paragraph_format.page_break_before
    dst.paragraph_format.widow_control = src.paragraph_format.widow_control


def insert_paragraph_after(paragraph, text: str = "", style_source=None, align=None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = paragraph._parent.add_paragraph()
    new_para._p = new_p
    if style_source is not None:
        clone_paragraph_format(style_source, new_para)
    if text:
        new_para.add_run(text)
    if align is not None:
        new_para.alignment = align
    return new_para


def find_anchor_paragraph(doc: Document, prefix: str):
    normalized_prefix = normalize(prefix)
    for idx, para in enumerate(doc.paragraphs):
        text = normalize(para.text)
        if text.startswith(normalized_prefix):
            return idx, para
    raise ValueError(f"未找到锚点标题: {prefix}")


def find_body_style_source(doc: Document, anchor_idx: int, anchor_para):
    paragraphs = doc.paragraphs
    start_idx = anchor_idx + 1
    for para in paragraphs[start_idx:]:
        text = para.text.strip()
        style_name = para.style.name if para.style else ""
        if text and not style_name.startswith("Heading"):
            return para
    return anchor_para


def insert_figure_sequence(doc: Document, fig: FigureSpec, cursor_para, style_source):
    intro_para = insert_paragraph_after(cursor_para, fig.intro_text, style_source=style_source)
    intro_para.alignment = style_source.alignment
    image_para = insert_paragraph_after(intro_para, "", style_source=style_source, align=WD_ALIGN_PARAGRAPH.CENTER)
    run = image_para.add_run()
    run.add_picture(str(fig.image_path), width=Inches(fig.width_inches))
    spacer = insert_paragraph_after(image_para, "", style_source=style_source)
    return spacer


def main():
    if not DOCX_PATH.exists():
        raise FileNotFoundError(DOCX_PATH)
    if not BACKUP_PATH.exists():
        shutil.copy2(DOCX_PATH, BACKUP_PATH)

    for fig in FIGURES:
        if not fig.image_path.exists():
            raise FileNotFoundError(f"图片不存在: {fig.image_path}")

    doc = Document(str(DOCX_PATH))

    anchor_cursors = {}
    anchor_style_sources = {}
    for fig in FIGURES:
        if fig.anchor_prefix not in anchor_cursors:
            anchor_idx, anchor_para = find_anchor_paragraph(doc, fig.anchor_prefix)
            anchor_cursors[fig.anchor_prefix] = anchor_para
            anchor_style_sources[fig.anchor_prefix] = find_body_style_source(doc, anchor_idx, anchor_para)

        cursor = anchor_cursors[fig.anchor_prefix]
        style_source = anchor_style_sources[fig.anchor_prefix]
        anchor_cursors[fig.anchor_prefix] = insert_figure_sequence(doc, fig, cursor, style_source)

    doc.save(str(DOCX_PATH))
    print(f"updated: {DOCX_PATH}")
    print(f"backup: {BACKUP_PATH}")


if __name__ == "__main__":
    main()
